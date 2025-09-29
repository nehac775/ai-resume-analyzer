"""
AI Resume Analyzer (no external spaCy model required)

Why this version?
-----------------
Streamlit Cloud was failing to load `en_core_web_sm`. To keep things simple,
I removed the dependency on downloaded spaCy models and now use spaCy's
lightweight English tokenizer + stopwords only (no large model).
This keeps the analyzer fast and avoids extra installation steps.

What it does
------------
- Upload a resume (PDF/DOCX/TXT)
- Paste a Job Description (JD)
- Clean & tokenize text (lowercase, alpha-only, stopword filter)
- Compute TF-IDF cosine similarity (match score)
- Highlight missing JD keywords
- Give suggestions
- Generate a downloadable PDF report

TODOs (future me)
-----------------
- Add better keyword extraction (e.g., RAKE/KeyBERT) if I expand deps
- Add extra metrics (Jaccard overlap, entity overlap)
- Batch mode: analyze multiple resumes against one JD
"""

import streamlit as st
import pandas as pd
import re
from io import BytesIO

# Light spaCy components (no external model download)
from spacy.lang.en import English
from spacy.lang.en.stop_words import STOP_WORDS

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Match your requirements.txt: using pypdf (not PyPDF2)
from pypdf import PdfReader
import docx
from fpdf import FPDF

# -------------------- Streamlit setup --------------------
st.set_page_config(page_title="AI Resume Analyzer", page_icon="📝", layout="wide")
st.title("📝 AI Resume Analyzer")
st.caption("Compares your resume to a job description, shows a match score, missing keywords, suggestions, and a downloadable report.")

# -------------------- NLP: lightweight tokenizer --------------------
# Using a blank English pipeline avoids downloading a model and still gives us a good tokenizer.
@st.cache_resource
def get_tokenizer():
    nlp = English()
    return nlp.tokenizer

TOKENIZER = get_tokenizer()
STOPWORDS = set(STOP_WORDS)

# -------------------- Helpers --------------------
def extract_text(file) -> str:
    """
    Extract raw text from PDF/DOCX/TXT uploads.
    I switched to pypdf (matches my requirements.txt) to avoid module mismatch.
    """
    name = (file.name or "").lower()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(file)
            pages = []
            for p in reader.pages:
                # Some PDFs return None; guard against that
                pages.append(p.extract_text() or "")
            return "\n".join(pages)
        elif name.endswith(".docx"):
            d = docx.Document(file)
            return "\n".join([p.text for p in d.paragraphs])
        else:
            # default: txt
            return file.read().decode("utf-8", errors="ignore")
    except Exception as e:
        # If extraction fails, surface the error and return empty text
        st.error(f"Could not read file: {e}")
        return ""

def clean_text(text: str) -> str:
    """
    Minimal cleaning:
    - lowercase
    - tokenize with spaCy's English tokenizer (no model)
    - keep alphabetic tokens that are not stopwords
    I used to lemmatize with a full model, but that requires downloads; TF-IDF is fine without it.
    """
    text = text.lower()
    doc = TOKENIZER(text)
    tokens = [t.text for t in doc if t.is_alpha and t.text not in STOPWORDS]
    return " ".join(tokens)

def compute_similarity(resume_clean: str, jd_clean: str) -> tuple[float, list[str]]:
    """
    TF-IDF cosine similarity between cleaned resume and JD.
    ngram_range=(1,2) slightly helps with short phrases like 'data analysis'.
    """
    vect = TfidfVectorizer(ngram_range=(1, 2), min_df=1)
    tfidf = vect.fit_transform([resume_clean, jd_clean])
    score = cosine_similarity(tfidf[0:1], tfidf[1:2])[0][0]
    return float(score), list(vect.get_feature_names_out())

def top_keywords(text: str, k: int = 30) -> list[str]:
    """
    Simple frequency-based keyword list from cleaned text.
    Good enough for highlighting gaps without extra deps.
    """
    words = re.findall(r"[a-z]+", text.lower())
    if not words:
        return []
    s = pd.Series(words)
    return s.value_counts().head(k).index.tolist()

def find_missing_keywords(resume_clean: str, jd_clean: str, limit: int = 20) -> list[str]:
    """
    JD keywords that aren't prominent in the resume.
    I limit the count so it doesn't overwhelm the user.
    """
    jd_kw = set(top_keywords(jd_clean, 50))
    res_kw = set(top_keywords(resume_clean, 50))
    missing = [w for w in jd_kw - res_kw if len(w) > 2]
    return sorted(missing)[:limit]

def generate_pdf_report(score: float, missing: list[str], suggestions: list[str]) -> bytes:
    """
    Tiny PDF report via FPDF. I keep styling simple (intern project vibes).
    """
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "AI Resume Analyzer Report", ln=True, align="C")

    pdf.ln(4)
    pdf.set_font("Arial", "", 12)
    pdf.cell(0, 8, f"Match Score: {score:.2%}", ln=True)

    pdf.ln(2)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Missing Keywords:", ln=True)
    pdf.set_font("Arial", "", 12)
    if missing:
        pdf.multi_cell(0, 6, ", ".join(missing))
    else:
        pdf.multi_cell(0, 6, "None 🎉")

    pdf.ln(2)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Suggestions:", ln=True)
    pdf.set_font("Arial", "", 12)
    if suggestions:
        for s in suggestions:
            pdf.multi_cell(0, 6, f"- {s}")
    else:
        pdf.multi_cell(0, 6, "Looks solid!")

    return pdf.output(dest="S").encode("latin1")

# -------------------- UI --------------------
col1, col2 = st.columns(2)
with col1:
    resume_file = st.file_uploader("Upload Resume (PDF, DOCX, or TXT)", type=["pdf", "docx", "txt"])
with col2:
    jd_text = st.text_area("Paste Job Description")

if resume_file and jd_text:
    resume_raw = extract_text(resume_file)
    if not resume_raw.strip():
        st.error("Could not read any text from the resume file.")
        st.stop()

    # Clean both texts using the lightweight tokenizer path
    resume_clean = clean_text(resume_raw)
    jd_clean     = clean_text(jd_text)

    # Compute similarity
    score, _ = compute_similarity(resume_clean, jd_clean)

    st.subheader("📊 Match Score")
    st.metric(label="Resume vs Job Description", value=f"{score:.2%}")

    # Missing keywords
    missing = find_missing_keywords(resume_clean, jd_clean)
    st.subheader("❌ Missing Keywords")
    st.write(", ".join(missing) if missing else "No major gaps found 🎉")

    # Suggestions (kept simple and honest)
    suggestions = []
    if score < 0.6:
        suggestions.append("Add more role-specific keywords from the job description.")
    if len(missing) > 0:
        suggestions.append("Weave missing skills naturally into bullet points (not a keyword dump).")
    if len(resume_clean.split()) < 120:
        suggestions.append("Your resume text seems short—expand impact bullets with metrics if possible.")
    if not suggestions:
        suggestions.append("Your resume already aligns well with this JD!")

    st.subheader("💡 Suggestions")
    for s in suggestions:
        st.write(f"- {s}")

    # Downloadable report
    pdf_bytes = generate_pdf_report(score, missing, suggestions)
    st.download_button("⬇️ Download Report (PDF)", data=pdf_bytes,
                       file_name="resume_analysis.pdf", mime="application/pdf")

else:
    st.info("👉 Upload a resume and paste a job description to analyze.")

